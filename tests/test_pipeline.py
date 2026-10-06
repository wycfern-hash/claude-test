import json

import pytest
from fastapi.testclient import TestClient

from shopee_clips import config, db


@pytest.fixture(autouse=True)
def tmp_data(tmp_path, monkeypatch):
    for name, sub in (("DATA_DIR", ""), ("DB_PATH", "clips.db"), ("REF_DIR", "ref"), ("IMG_DIR", "images"),
                      ("VID_DIR", "videos"), ("BROWSER_PROFILE", "bp")):
        monkeypatch.setattr(config, name, tmp_path / sub if sub else tmp_path)


def test_key_parsing():
    assert db.shopee_key("https://shopee.tw/Foo-Bar-i.123.456?sp=1") == "123.456"
    assert db.shopee_key("https://shopee.tw/product/123/456") == "123.456"
    assert db.shopee_key("https://s.shopee.tw/abc") is None


def test_one_product_one_video():
    with db.connect() as conn:
        assert db.add_product(conn, "https://shopee.tw/A-i.1.2")
        assert db.add_product(conn, "https://shopee.tw/other-name-i.1.2?x=y") is None
        with pytest.raises(ValueError):
            db.add_product(conn, "https://s.shopee.tw/short")


def test_state_machine():
    with db.connect() as conn:
        pid = db.add_product(conn, "https://shopee.tw/A-i.1.2")
        with pytest.raises(ValueError):
            db.move(conn, pid, "uploaded")  # 不能跳過審核
        db.move(conn, pid, "image_review")
        db.move(conn, pid, "image_approved", selected_image="images/1/0.png")
        db.move(conn, pid, "video_review", video_path="videos/1/final.mp4")
        db.move(conn, pid, "video_approved")
        db.move(conn, pid, "uploaded", uploaded_at=db.now())
        assert db.uploaded_today(conn) == 1


def test_web_review_flow():
    from shopee_clips.web import app

    with db.connect() as conn:
        pid = db.add_product(conn, "https://shopee.tw/A-i.1.2", title="測試")
        db.move(conn, pid, "image_review", images=json.dumps(["images/1/0.png"]))
    c = TestClient(app)
    assert "測試" in c.get("/images").text
    c.post(f"/images/{pid}", data={"act": "approve", "selected": "images/1/0.png"})
    with db.connect() as conn:
        assert db.get(conn, pid)["status"] == "image_approved"
        db.move(conn, pid, "video_review", video_path="videos/1/final.mp4", video_title="t", video_caption="c")
    c.post(f"/videos/{pid}", data={"act": "approve", "video_title": "改過", "video_caption": "c"})
    with db.connect() as conn:
        r = db.get(conn, pid)
        assert (r["status"], r["video_title"]) == ("video_approved", "改過")
    assert "下載影片" in c.get("/ready").text
    c.post(f"/ready/{pid}")
    with db.connect() as conn:
        assert db.get(conn, pid)["status"] == "uploaded"


def test_concat_trim(tmp_path):
    import shutil
    import subprocess

    from shopee_clips.videogen import concat_trim

    if not shutil.which("ffmpeg"):
        pytest.skip("no ffmpeg")
    clips = []
    for i in range(2):
        p = tmp_path / f"c{i}.mp4"
        subprocess.run(["ffmpeg", "-y", "-f", "lavfi", "-i", "color=c=red:s=180x320:d=8:r=15", "-pix_fmt", "yuv420p", str(p)],
                       check=True, capture_output=True)
        clips.append(p)
    out = tmp_path / "final.mp4"
    concat_trim(clips, out)
    dur = subprocess.run(["ffprobe", "-v", "error", "-show_entries", "format=duration", "-of", "csv=p=0", str(out)],
                         capture_output=True, text=True).stdout
    assert 14.5 < float(dur) < 15.5


def test_template_script_requires_points_and_never_invents():
    from shopee_clips import scriptgen

    with db.connect() as conn:
        pid = db.add_product(conn, "https://shopee.tw/A-i.1.2", title="保溫杯")
        with pytest.raises(RuntimeError):
            scriptgen.from_template(db.get(conn, pid))
        db.update(conn, pid, script=json.dumps({"user_points": ["保溫12小時", "304不鏽鋼", "一鍵開蓋"]}))
        s = scriptgen.get_script(db.get(conn, pid))
        assert s["selling_points"] == ["保溫12小時", "304不鏽鋼", "一鍵開蓋"] and "保溫杯" in s["hook"]


def test_manual_upload_to_slideshow_to_video_review(monkeypatch):
    """免費路徑：上傳 5 張圖 → 自動合成 15 秒影片 → 進審片。"""
    import io
    import shutil
    import subprocess

    from PIL import Image

    from shopee_clips import videogen
    from shopee_clips.web import app

    if not shutil.which("ffmpeg"):
        pytest.skip("no ffmpeg")
    monkeypatch.setattr(config, "VIDEO_MODE", "slideshow")
    monkeypatch.setattr(config, "TTS", False)
    with db.connect() as conn:
        pid = db.add_product(conn, "https://shopee.tw/A-i.1.2", title="保溫杯")
    files = []
    for i in range(5):
        buf = io.BytesIO()
        Image.new("RGB", (800, 1000), (40 * i, 100, 200)).save(buf, "PNG")
        files.append(("files", (f"{i}.png", buf.getvalue(), "image/png")))
    c = TestClient(app)
    c.post(f"/todo/{pid}", files=files, data={"points": "保溫12小時\n304不鏽鋼\n一鍵開蓋"})
    with db.connect() as conn:
        assert db.get(conn, pid)["status"] == "image_approved"
        assert videogen.run(conn) == 1
        r = db.get(conn, pid)
        assert r["status"] == "video_review" and "保溫12小時" in r["video_caption"]
        out = config.DATA_DIR / r["video_path"]
    dur = subprocess.run(["ffprobe", "-v", "error", "-show_entries", "format=duration", "-of", "csv=p=0", str(out)],
                         capture_output=True, text=True).stdout
    assert 14.5 < float(dur) < 15.5


def test_flow_clips_to_15s_review(monkeypatch):
    """Flow 路徑：上傳兩段 8 秒 mp4 → 接起來裁 15 秒、直式 1080x1920 → 進審片。"""
    import shutil
    import subprocess

    from shopee_clips import videogen

    if not shutil.which("ffmpeg"):
        pytest.skip("no ffmpeg")
    monkeypatch.setattr(config, "TTS", False)
    d = config.DATA_DIR
    clips = []
    for i in range(2):
        p = d / f"in{i}.mp4"
        subprocess.run(["ffmpeg", "-y", "-f", "lavfi", "-i", "testsrc=s=1280x720:d=8:r=24", "-f", "lavfi", "-i",
                        "sine=d=8", "-pix_fmt", "yuv420p", "-shortest", str(p)], check=True, capture_output=True)
        clips.append(p)
    with db.connect() as conn:
        pid = db.add_product(conn, "https://shopee.tw/A-i.1.2", title="保溫杯")
        db.update(conn, pid, script=json.dumps({"user_points": ["a", "b", "c"]}))
        db.move(conn, pid, "image_approved", selected_image="images/1/0.png")
        videogen.finish_flow(conn, pid, clips)
        r = db.get(conn, pid)
        assert r["status"] == "video_review"
        out = config.DATA_DIR / r["video_path"]
    probe = subprocess.run(["ffprobe", "-v", "error", "-show_entries", "stream=width,height:format=duration",
                            "-of", "csv=p=0", str(out)], capture_output=True, text=True).stdout
    assert "1080,1920" in probe and any(14.5 < float(x) < 15.5 for x in probe.split() if x.replace(".", "").isdigit())


def test_no_label_text_in_captions():
    from shopee_clips import slideshow

    texts = slideshow.beats({"hook": "開箱保溫杯！", "selling_points": ["a", "b", "c"], "cta": "點連結"})
    assert not any(w in t for t in texts for w in ("hook", "賣點", "CTA"))


def test_all_pages_render():
    from shopee_clips.web import app

    c = TestClient(app)
    for path in ("/", "/todo", "/images", "/flow", "/videos", "/ready"):
        assert c.get(path).status_code == 200, path


def _make_xlsx(path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["序號", "商品名稱", "商品連結", "價格", "賣點1", "賣點2"])
    ws.append([1, "保溫杯", "https://shopee.tw/保溫杯-i.11.22", 399, "保溫12小時", "304不鏽鋼"])
    ws.append([2, "筆", "https://shopee.tw/pen-i.33.44", 20, "", ""])
    ws.append([3, "重複", "https://shopee.tw/other-i.11.22?x=1", 1, "", ""])          # 同商品 → 去重
    ws.append([4, "短連結", "https://s.shopee.tw/abc", 1, "", ""])                    # 需展開
    ws.append([5, "超連結", "點我", 1, "", ""])
    ws.cell(row=6, column=3).hyperlink = "https://shopee.tw/x-i.55.66"
    ws.append([6, "壞資料", "https://s.shopee.tw/bad", 1, "", ""])
    wb.save(path)


def test_import_excel(tmp_path, monkeypatch):
    from shopee_clips import sourcing

    def fake_expand(url):
        if url.endswith("/abc"):
            return "https://shopee.tw/product/77/88"
        raise RuntimeError("redirect failed")

    monkeypatch.setattr(sourcing, "expand_short", fake_expand)
    x = tmp_path / "選品.xlsx"
    _make_xlsx(x)
    with db.connect() as conn:
        r = sourcing.import_excel(conn, str(x))
        assert (r["added"], r["dup"], len(r["failed"])) == (4, 1, 1)
        keys = {row["shopee_key"] for row in conn.execute("SELECT shopee_key FROM products")}
        assert keys == {"11.22", "33.44", "77.88", "55.66"}
        cup = conn.execute("SELECT * FROM products WHERE shopee_key='11.22'").fetchone()
        assert cup["title"] == "保溫杯" and json.loads(cup["script"])["user_points"] == ["保溫12小時", "304不鏽鋼"]


def test_excel_without_url_column(tmp_path):
    from openpyxl import Workbook

    from shopee_clips import sourcing

    wb = Workbook()
    wb.active.append(["商品名稱", "價格"])
    wb.save(tmp_path / "a.xlsx")
    with db.connect() as conn:
        r = sourcing.import_excel(conn, str(tmp_path / "a.xlsx"))
    assert r["added"] == 0 and r["failed"]


def test_site_labels_are_valid_regex():
    import re

    from shopee_clips.webauto import site

    for name in ("gemini", "flow"):
        for k, v in site(name).items():
            if isinstance(v, str) and not k.startswith(("url", "prompt_prefix")):
                re.compile(v)
